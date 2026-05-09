import os
import sys
import json
import datetime
import argparse
import logging
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[logging.StreamHandler(sys.stdout)]
)
logger = logging.getLogger(__name__)

# 配置路径
SAMPLE_DOCX_PATH = os.path.join(os.path.dirname(__file__), '..', 'assets', 'sample.docx')

# 检查样本文档是否存在
if not os.path.exists(SAMPLE_DOCX_PATH):
    logger.error(f"Error: Sample document not found at {SAMPLE_DOCX_PATH}")
    sys.exit(1)

def update_docx(json_data, output_dir: str):
    """更新sample.docx文件"""
    # 从JSON数据中提取信息
    english_info = json_data.get('english', {})

    # 确保输出目录存在
    os.makedirs(output_dir, exist_ok=True)

    # 基于现有样本文档创建，保留页眉页脚
    doc = Document(SAMPLE_DOCX_PATH)

    # 清空文档内容，保留页眉页脚
    for paragraph in doc.paragraphs:
        p = paragraph._element
        p.getparent().remove(p)
        p._p = p._element = None

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    # 设置东亚字体
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    # 添加内容
    # Paragraph 1: RADIOLOGY REPORT
    p1 = doc.add_paragraph()
    run1 = p1.add_run('RADIOLOGY REPORT')
    run1.bold = True
    run1.font.size = Pt(16)  # 三号字体
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 2: MRI Type
    p2 = doc.add_paragraph()
    run2 = p2.add_run(f"{english_info.get('EXAMINATION TYPE', '<无识别结果请补充>')}")
    run2.bold = True
    run2.font.size = Pt(16)  # 三号字体
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 空行
    doc.add_paragraph()

    # Paragraph 3-9: Patient Information
    p3 = doc.add_paragraph(f"NAME: {english_info.get('NAME', '<无识别结果请补充>')}")
    p4 = doc.add_paragraph(f"MRN: {english_info.get('MRN', '<无识别结果请补充>')}")
    p5 = doc.add_paragraph(f"DOB: {english_info.get('DOB', '<无识别结果请补充>')}")
    p6 = doc.add_paragraph(f"GENDER: {english_info.get('GENDER', '<无识别结果请补充>')}")
    p7 = doc.add_paragraph(f"AGE: {english_info.get('AGE', '<无识别结果请补充>')}")
    p8 = doc.add_paragraph(f"RADIOLOGY NO: {english_info.get('RADIOLOGY NO', '<无识别结果请补充>')}")
    p9 = doc.add_paragraph(f"DATE: {english_info.get('DATE', '<无识别结果请补充>')}")

    # Empty paragraph
    doc.add_paragraph()

    # Paragraph 11: Examination Type
    p11 = doc.add_paragraph()
    run = p11.add_run(f"{english_info.get('EXAMINATION SITE', '<无识别结果请补充>')}-{english_info.get('EXAMINATION TYPE', '<无识别结果请补充>')}")
    run.bold = True
    run.underline = True

    # Paragraph 12: Technique
    p12 = doc.add_paragraph()
    p12.add_run('Technique: ').bold = True
    p12.add_run(english_info.get('Technique', '<无识别结果请补充>'))

    # Empty paragraph
    doc.add_paragraph()

    # Paragraph 14: Findings
    p14 = doc.add_paragraph()
    p14.add_run('Findings: ').bold = True
    p14.add_run(english_info.get('Findings', '<无识别结果请补充>'))

    # Empty paragraph
    doc.add_paragraph()

    # Paragraph 16: Impressions header
    p16 = doc.add_paragraph()
    p16.add_run('Impressions:').bold = True

    # Paragraph 17-18: Impressions list
    impressions = english_info.get('Impressions', '<无识别结果请补充>')
    if impressions:
        if isinstance(impressions, list):
            items = impressions
        else:
            items = [item.strip() for item in impressions.split(';') if item.strip()]

        if len(items) <= 1:
            for item in items:
                p = doc.add_paragraph(item)
        else:
            for item in items:
                p = doc.add_paragraph()
                p.add_run(f"•{item}")
    else:
        p17 = doc.add_paragraph('No significant findings.')

    # Empty paragraph
    doc.add_paragraph()

    # Paragraph 20-23: Footer info
    p20 = doc.add_paragraph(f"Reported by: {english_info.get('REPORTING PHYSICIAN', '<无识别结果请补充>')}")
    p21 = doc.add_paragraph(f"Reported date: {english_info.get('DATE', '<无识别结果请补充>')}")
    p22 = doc.add_paragraph('Translated by: DR LISA KHOO')
    p23 = doc.add_paragraph('Issued by Shanghai East Hospital')

    # 生成文件名：MRN_NAME_当前时间.docx
    mrn = english_info.get('MRN', '<无识别结果请补充>')
    name = english_info.get('NAME', '<无识别结果请补充>').replace(' ', '_')
    current_time = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"{mrn}_{name}_{current_time}.docx"
    output_path = os.path.join(output_dir, filename)

    # 保存文档
    doc.save(output_path)
    logger.info(f"Updated document saved at {output_path}")

def main():
    """主函数"""
    parser = argparse.ArgumentParser(description="Generate hospital report Word document")
    parser.add_argument(
        "--json-file",
        required=True,
        help="Path to JSON file containing report data",
    )
    parser.add_argument(
        "--output-dir",
        required=True,
        help="Output directory for generated Word document",
    )

    args = parser.parse_args()

    json_file_path = os.path.abspath(args.json_file)
    logger.info(f"Processing JSON file: {json_file_path}")

    # 检查文件是否存在
    if not os.path.exists(json_file_path):
        logger.error(f"Error: File '{json_file_path}' not found")
        sys.exit(1)

    # 读取JSON文件
    try:
        with open(json_file_path, 'r', encoding='utf-8') as f:
            json_data = json.load(f)
        logger.info("Loaded JSON data:")
        logger.info(json.dumps(json_data, ensure_ascii=False, indent=2))
    except Exception as e:
        logger.error(f"Error reading JSON file: {e}")
        sys.exit(1)

    # 更新文档
    update_docx(json_data, args.output_dir)
    logger.info("Process completed successfully!")

if __name__ == "__main__":
    main()
