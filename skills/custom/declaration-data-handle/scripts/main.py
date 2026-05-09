#!/usr/bin/env python3
import os
import sys
import json
import argparse
import uuid
import re
from pathlib import Path
from datetime import datetime
from collections import defaultdict

os.environ.setdefault('HF_ENDPOINT', 'https://hf-mirror.com')

from fuzzywuzzy import fuzz

# 检查 openpyxl 是否可用
try:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# 检查 lanceDB 是否可用
try:
    import lancedb
    from sentence_transformers import SentenceTransformer
    LANCEDB_AVAILABLE = True
    _embedding_model = None
except ImportError:
    LANCEDB_AVAILABLE = False

# 正则表达式模式
PLACEHOLDER_PATTERN = re.compile(r'\{([^}]+)\}')
MULTI_PIPE_PATTERN = re.compile(r'\|+')
PIPE_START_PATTERN = re.compile(r'^\|+')
PIPE_END_PATTERN = re.compile(r'\|+$')
MULTI_SLASH_PATTERN = re.compile(r'/+')

# 字段路径映射表
FIELD_PATH_MAP = {
    "IEFlag": "DecMessage.DecHead.IEFlag",
    "BillNo": "DecMessage.DecHead.BillNo",
    "ContrNo": "DecMessage.DecHead.ContrNo",
    "CutMode": "DecMessage.DecHead.CutMode",
    "DistinatePort": "DecMessage.DecHead.DistinatePort",
    "FeeCurr": "DecMessage.DecHead.FeeCurr",
    "FeeMark": "DecMessage.DecHead.FeeMark",
    "FeeRate": "DecMessage.DecHead.FeeRate",
    "GrossWet": "DecMessage.DecHead.GrossWet",
    "InsurCurr": "DecMessage.DecHead.InsurCurr",
    "InsurMark": "DecMessage.DecHead.InsurMark",
    "InsurRate": "DecMessage.DecHead.InsurRate",
    "ManualNo": "DecMessage.DecHead.ManualNo",
    "NetWt": "DecMessage.DecHead.NetWt",
    "NoteS": "DecMessage.DecHead.NoteS",
    "OtherCurr": "DecMessage.DecHead.OtherCurr",
    "OtherMark": "DecMessage.DecHead.OtherMark",
    "OtherRate": "DecMessage.DecHead.OtherRate",
    "OwnerCode": "DecMessage.DecHead.OwnerCode",
    "OwnerCiqCode": "DecMessage.DecHead.OwnerCiqCode",
    "OwnerName": "DecMessage.DecHead.OwnerName",
    "PackNo": "DecMessage.DecHead.PackNo",
    "TradeCountry": "DecMessage.DecHead.TradeCountry",
    "TradeMode": "DecMessage.DecHead.TradeMode",
    "TrafMode": "DecMessage.DecHead.TrafMode",
    "TrafName": "DecMessage.DecHead.TrafName",
    "VoyNo": "DecMessage.DecHead.VoyNo",
    "TransMode": "DecMessage.DecHead.TransMode",
    "WrapType": "DecMessage.DecHead.WrapType",
    "TradeArea": "DecMessage.DecHead.TradeArea",
    "DespPort": "DecMessage.DecHead.DespPort",
    "OverseasConsignorAEOCode": "DecMessage.DecHead.OverseasConsignorAEOCode",
    "OverseasConsigneeEname": "DecMessage.DecHead.OverseasConsigneeEname",
    "TradeCode": "DecMessage.DecHead.TradeCode",
    "TradeCiqCode": "DecMessage.DecHead.TradeCiqCode",
    "TradeName": "DecMessage.DecHead.TradeName",
    "TradeNameEn": "DecMessage.DecHead.TradeNameEn"
}


def reorder_results_with_keyword_match(results, search_text, text_field_names):
    """
    根据关键词匹配重新排序搜索结果
    
    Args:
        results: 搜索结果
        search_text: 搜索文本
        text_field_names: 文本字段名称列表
    
    Returns:
        重新排序后的结果
    """
    if not search_text or not results.get('ids') or len(results['ids'][0]) == 0:
        return results
    
    ids = results['ids'][0]
    docs = results['documents'][0]
    metas = results['metadatas'][0] if results.get('metadatas') else [None] * len(ids)
    distances = results['distances'][0] if results.get('distances') else [None] * len(ids)
    
    search_terms = search_text.lower().split()
    
    def count_matching_keywords(text):
        text_lower = text.lower()
        return sum(1 for term in search_terms if term in text_lower)
    
    scored_results = []
    for i in range(len(ids)):
        meta = metas[i]
        doc = docs[i]
        
        meta_texts = []
        if meta:
            for field_name in text_field_names:
                field_value = meta.get(field_name, '')
                if field_value:
                    meta_texts.append(str(field_value))
        meta_text = ' '.join(meta_texts)
        
        score = 0
        search_lower = search_text.lower()
        meta_lower = meta_text.lower()
        doc_lower = doc.lower()
        
        if meta_lower == search_lower:
            score = 1000
        elif search_lower in meta_lower:
            score = 500
        elif search_lower in doc_lower:
            score = 200
        else:
            keyword_count = count_matching_keywords(meta_text + ' ' + doc)
            score = keyword_count * 10
        
        scored_results.append((-score, distances[i] if distances[i] is not None else float('inf'), ids[i], docs[i], metas[i], distances[i]))
    
    scored_results.sort()
    
    new_ids = [x[2] for x in scored_results]
    new_docs = [x[3] for x in scored_results]
    new_metas = [x[4] for x in scored_results]
    new_distances = [x[5] for x in scored_results]
    
    return {
        'ids': [new_ids],
        'documents': [new_docs],
        'metadatas': [new_metas] if new_metas[0] is not None else None,
        'distances': [new_distances] if new_distances[0] is not None else None
    }


def get_embedding_model():
    """
    获取嵌入模型（单例模式）
    
    Returns:
        嵌入模型实例
    """
    global _embedding_model
    if _embedding_model is None:
        _embedding_model = SentenceTransformer('all-MiniLM-L6-v2')
    return _embedding_model


def query_from_db(persist_directory, collection_name, query_text=None, filter_conditions=None, 
                  n_results=10, strict_match_fields=None):
    """
    从向量数据库查询数据
    
    Args:
        persist_directory: 数据库持久化目录
        collection_name: 集合名称
        query_text: 查询文本
        filter_conditions: 过滤条件
        n_results: 返回结果数量
        strict_match_fields: 需要严格匹配的字段
    
    Returns:
        查询结果
    """
    db = lancedb.connect(persist_directory)
    
    try:
        table = db.open_table(collection_name)
    except Exception as e:
        return None
    
    if filter_conditions and len(filter_conditions) > 0:
        all_items = table.to_pandas()
        total_count = len(all_items)
        
        filtered_items = []
        
        for idx, (_, row) in enumerate(all_items.iterrows()):
            match = True
            for key, value in filter_conditions.items():
                meta_value = str(row.get(key, ''))
                
                if strict_match_fields and key in strict_match_fields:
                    if not meta_value or (str(value).lower() not in meta_value.lower() and meta_value.lower() not in str(value).lower()):
                        match = False
                        break
            
            if match:
                filtered_items.append(row)
        
        if filtered_items:
            if query_text:
                model = get_embedding_model()
                query_vector = model.encode(query_text).tolist()
                
                results_df = table.search(query_vector).limit(n_results).to_pandas()
                
                if strict_match_fields and filter_conditions:
                    mask = []
                    for _, row in results_df.iterrows():
                        item_match = True
                        for key, value in filter_conditions.items():
                            if key in strict_match_fields:
                                meta_value = str(row.get(key, ''))
                                if not meta_value or (str(value).lower() not in meta_value.lower() and meta_value.lower() not in str(value).lower()):
                                    item_match = False
                                    break
                        mask.append(item_match)
                    results_df = results_df[mask]
                
                ids = results_df['id'].tolist()
                documents = results_df['document'].tolist()
                metadatas = []
                for _, row in results_df.iterrows():
                    meta = row.drop(['id', 'document', 'vector']).to_dict()
                    metadatas.append(meta)
                distances = results_df['_distance'].tolist() if '_distance' in results_df.columns else None
                
                results = {
                    'ids': [ids],
                    'documents': [documents],
                    'metadatas': [metadatas],
                    'distances': [distances] if distances else None
                }
            else:
                ids = [item['id'] for item in filtered_items]
                documents = [item['document'] for item in filtered_items]
                metadatas = []
                for item in filtered_items:
                    meta = {k: v for k, v in item.items() if k not in ['id', 'document', 'vector']}
                    metadatas.append(meta)
                
                results = {
                    'ids': [ids],
                    'documents': [documents],
                    'metadatas': [metadatas],
                    'distances': None
                }
        else:
            if query_text:
                model = get_embedding_model()
                query_vector = model.encode(query_text).tolist()
                results_df = table.search(query_vector).limit(n_results).to_pandas()
                
                ids = results_df['id'].tolist()
                documents = results_df['document'].tolist()
                metadatas = []
                for _, row in results_df.iterrows():
                    meta = row.drop(['id', 'document', 'vector']).to_dict()
                    metadatas.append(meta)
                distances = results_df['_distance'].tolist() if '_distance' in results_df.columns else None
                
                results = {
                    'ids': [ids],
                    'documents': [documents],
                    'metadatas': [metadatas],
                    'distances': [distances] if distances else None
                }
            else:
                results = None
    else:
        if query_text:
            model = get_embedding_model()
            query_vector = model.encode(query_text).tolist()
            results_df = table.search(query_vector).limit(n_results).to_pandas()
            
            ids = results_df['id'].tolist()
            documents = results_df['document'].tolist()
            metadatas = []
            for _, row in results_df.iterrows():
                meta = row.drop(['id', 'document', 'vector']).to_dict()
                metadatas.append(meta)
            distances = results_df['_distance'].tolist() if '_distance' in results_df.columns else None
            
            results = {
                'ids': [ids],
                'documents': [documents],
                'metadatas': [metadatas],
                'distances': [distances] if distances else None
            }
        else:
            results = None
    
    return results


def load_analysis_results(analysis_file):
    """
    加载分析结果 JSON 文件
    
    Args:
        analysis_file: 文件路径
    
    Returns:
        分析结果数据或 None
    """
    if analysis_file and os.path.exists(analysis_file):
        try:
            with open(analysis_file, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            return None
    else:
        return None


def merge_json_objects_with_guid(json_list_with_guid):
    """
    合并包含 GUID 的 JSON 对象列表
    
    Args:
        json_list_with_guid: JSON 对象列表，每个对象包含 guid 和 analysis
    
    Returns:
        (合并后的数据, GUID 到业务类型的映射)
    """
    if not json_list_with_guid:
        return {}, {}

    invoice_data = None
    default_data = None
    other_data_list = []

    # 分离不同类型的单证
    for idx, data_with_guid in enumerate(json_list_with_guid):
        json_data = data_with_guid.get('analysis', {}) or data_with_guid
        file_guid = data_with_guid.get('guid', str(uuid.uuid4()))
        file_business_type = json_data.get('FileBusinessType', '')
        file_name = json_data.get('FileName', f'未知文件_{idx+1}')

        if file_business_type == '发票':
            invoice_data = {'json': json_data, 'guid': file_guid}
        elif file_business_type == '企业默认值':
            default_data = {'json': json_data, 'guid': file_guid}
        else:
            other_data_list.append({'json': json_data, 'guid': file_guid})

    # 按商品描述分组
    dec_list_groups = {}
    existing_items = []
    group_invoice_gno = {}

    # 首先处理发票数据作为基准
    if invoice_data:
        json_data = invoice_data['json']
        file_guid = invoice_data['guid']

        if 'DecMessage' in json_data and 'DecLists' in json_data['DecMessage']:
            dec_list = json_data['DecMessage']['DecLists'].get('DecList', [])
            if isinstance(dec_list, list):
                for item_idx, item in enumerate(dec_list):
                    if isinstance(item, dict):
                        description = item.get('Description', '')
                        qty1 = item.get('Qty1')
                        gno = item.get('GNo')
                        if description:
                            qty1_str = str(qty1) if qty1 is not None else 'None'
                            group_key = f"{description}|{qty1_str}"
                            existing_items.append({
                                'desc': description,
                                'qty1': qty1,
                                'key': group_key
                            })
                            group_invoice_gno[group_key] = gno if gno is not None else float('inf')

                            if group_key not in dec_list_groups:
                                dec_list_groups[group_key] = []
                            dec_list_groups[group_key].append({
                                'guid': file_guid,
                                'item': item
                            })

    # 处理其他单证，通过模糊匹配合并到对应组
    for data_idx, data_with_guid in enumerate(other_data_list):
        json_data = data_with_guid['json']
        file_guid = data_with_guid['guid']

        if 'DecMessage' in json_data and 'DecLists' in json_data['DecMessage']:
            dec_list = json_data['DecMessage']['DecLists'].get('DecList', [])
            if isinstance(dec_list, list):
                for item_idx, item in enumerate(dec_list):
                    if isinstance(item, dict):
                        description = item.get('Description', '')
                        if description:
                            best_match, score = find_best_match_by_description(
                                description, existing_items, threshold=70
                            )

                            if best_match:
                                group_key = best_match
                            elif not invoice_data:
                                qty1 = item.get('Qty1')
                                qty1_str = str(qty1) if qty1 is not None else 'None'
                                group_key = f"{description}|{qty1_str}"
                                existing_items.append({
                                    'desc': description,
                                    'qty1': qty1,
                                    'key': group_key
                                })
                            else:
                                continue

                            if group_key not in dec_list_groups:
                                dec_list_groups[group_key] = []
                            dec_list_groups[group_key].append({
                                'guid': file_guid,
                                'item': item
                            })

    result = {}
    guid_to_business_type = {}

    # 递归合并字段
    def merge_recursive(target, source, source_guid, path="root"):
        if isinstance(source, dict):
            for key, value in source.items():
                current_path = f"{path}.{key}"
                if key == 'DecLists':
                    continue

                if key == 'FileBusinessType' and source_guid:
                    guid_to_business_type[source_guid] = value
                    if key not in target and value and value != '':
                        target[key] = value
                    continue

                if key not in target:
                    target[key] = {} if isinstance(value, (dict, list)) else []

                if isinstance(value, dict):
                    if not isinstance(target[key], dict):
                        target[key] = {}
                    merge_recursive(target[key], value, source_guid, current_path)
                elif isinstance(value, list):
                    if not isinstance(target[key], list):
                        target[key] = []
                    for idx_list, list_item in enumerate(value):
                        list_path = f"{current_path}[{idx_list}]"
                        if isinstance(list_item, dict):
                            found = False
                            for existing_item in target[key]:
                                if isinstance(existing_item, dict):
                                    merge_recursive(existing_item, list_item, source_guid, list_path)
                                    found = True
                                    break
                            if not found:
                                new_item = {}
                                merge_recursive(new_item, list_item, source_guid, list_path)
                                target[key].append(new_item)
                        else:
                            entry = {'guid': source_guid, 'value': list_item}
                            if entry not in target[key]:
                                target[key].append(entry)
                else:
                    if not isinstance(target[key], list):
                        target[key] = []
                    entry = {'guid': source_guid, 'value': value}
                    if entry not in target[key]:
                        target[key].append(entry)

    # 合并所有单证的非列表字段
    for data_idx, data_with_guid in enumerate(json_list_with_guid):
        json_data = data_with_guid.get('analysis', {}) or data_with_guid
        file_guid = data_with_guid.get('guid', str(uuid.uuid4()))
        merge_recursive(result, json_data, file_guid)

    # 合并商品列表
    if dec_list_groups:
        if 'DecMessage' not in result:
            result['DecMessage'] = {}
        if 'DecLists' not in result['DecMessage']:
            result['DecMessage']['DecLists'] = {}

        merged_dec_list = []
        sorted_groups = sorted(
            dec_list_groups.items(),
            key=lambda x: group_invoice_gno.get(x[0], float('inf'))
        )

        for goods_key, items_with_guid in sorted_groups:
            if items_with_guid:
                first_item = items_with_guid[0]['item']
                field_order = list(first_item.keys())
            else:
                field_order = []

            merged_item = {}
            for item_with_guid in items_with_guid:
                item_guid = item_with_guid['guid']
                item_data = item_with_guid['item']

                for key, value in item_data.items():
                    if key not in merged_item:
                        merged_item[key] = []

                    if isinstance(value, (dict, list)):
                        entry = {'guid': item_guid, 'value': value}
                        if entry not in merged_item[key]:
                            merged_item[key].append(entry)
                    else:
                        entry = {'guid': item_guid, 'value': value}
                        if entry not in merged_item[key]:
                            merged_item[key].append(entry)

            ordered_item = {}
            for key in field_order:
                if key in merged_item:
                    ordered_item[key] = merged_item[key]
            for key in merged_item:
                if key not in ordered_item:
                    ordered_item[key] = merged_item[key]

            merged_dec_list.append(ordered_item)

        result['DecMessage']['DecLists']['DecList'] = merged_dec_list

        # 应用企业默认值
        if default_data:
            default_json = default_data['json']
            default_guid = default_data['guid']
            if 'DecMessage' in default_json and 'DecLists' in default_json['DecMessage']:
                default_dec_list = default_json['DecMessage']['DecLists'].get('DecList', [])
                if isinstance(default_dec_list, list) and len(default_dec_list) > 0:
                    default_item = default_dec_list[0]

                    for item in result['DecMessage']['DecLists']['DecList']:
                        for key, value in default_item.items():
                            if key == 'GNo':
                                continue
                            if key not in item or len(item[key]) == 0:
                                item[key] = []
                                entry = {'guid': default_guid, 'value': value}
                                item[key].append(entry)

    # 过滤空值
    def filter_values(obj, parent_key=None):
        if isinstance(obj, dict):
            for key, value in obj.items():
                obj[key] = filter_values(value, key)
        elif isinstance(obj, list):
            if parent_key == 'FileBusinessType':
                return obj
            if len(obj) > 0 and isinstance(obj[0], dict) and 'guid' in obj[0] and 'value' in obj[0]:
                all_values = [item['value'] for item in obj]

                all_empty_or_zero = True
                for val in all_values:
                    if val is not None and val != '' and not (isinstance(val, (int, float)) and val == 0):
                        all_empty_or_zero = False
                        break

                if all_empty_or_zero:
                    return []

                filtered = []
                for item in obj:
                    if should_include_value(item['value'], all_values):
                        filtered.append(item)
                return filtered
            else:
                return [filter_values(item) for item in obj]
        return obj

    result = filter_values(result)
    result['FileBusinessType'] = '正式报关单'

    return result, guid_to_business_type


def find_best_match_by_description(new_desc, existing_items, threshold=70):
    """
    根据商品描述查找最佳匹配
    
    Args:
        new_desc: 新描述
        existing_items: 现有项列表
        threshold: 相似度阈值
    
    Returns:
        (最佳匹配项的key, 相似度分数)
    """
    best_match = None
    best_score = 0

    for existing_item in existing_items:
        score = calculate_description_similarity(new_desc, existing_item['desc'])
        if score > best_score:
            best_score = score
            best_match = existing_item['key']

    if best_score >= threshold:
        return best_match, best_score

    return None, best_score


def calculate_description_similarity(s1, s2):
    """
    计算两个描述的相似度
    
    Args:
        s1: 字符串1
        s2: 字符串2
    
    Returns:
        相似度分数（0-100）
    """
    if not s1 or not s2:
        return 0

    score1 = fuzz.token_sort_ratio(s1, s2)
    score2 = fuzz.partial_ratio(s1, s2)
    score3 = fuzz.ratio(s1, s2)

    return max(score1, score2, score3)


def should_include_value(value, all_values):
    """
    判断是否应该包含某个值
    
    Args:
        value: 待判断的值
        all_values: 所有值
    
    Returns:
        是否包含
    """
    has_non_empty = False
    has_non_zero = False

    for val in all_values:
        if val is not None and val != '':
            has_non_empty = True
        if isinstance(val, (int, float)) and val != 0:
            has_non_zero = True

    if value is None or value == '':
        return not has_non_empty
    if isinstance(value, (int, float)) and value == 0:
        return not has_non_zero

    return True


def get_value_from_merged_data(merged_data, field_path):
    """
    从合并后的数据中获取指定路径的值
    
    Args:
        merged_data: 合并后的数据
        field_path: 字段路径（用点分隔）
    
    Returns:
        字段值列表，每个元素包含 guid 和 value
    """
    parts = field_path.split('.')
    current = merged_data

    for part in parts:
        if isinstance(current, dict) and part in current:
            current = current[part]
        else:
            return []

    if isinstance(current, list) and len(current) > 0 and isinstance(current[0], dict) and 'guid' in current[0] and 'value' in current[0]:
        return current

    return [{'guid': None, 'value': current}]


def query_product_info(manufacturer=None, english_desc=None, script_dir=None):
    """
    查询产品信息
    
    Args:
        manufacturer: 制造商
        english_desc: 英文描述
        script_dir: 脚本目录
    
    Returns:
        产品信息字典或 None
    """
    if not LANCEDB_AVAILABLE:
        return None

    if not script_dir:
        script_dir = os.path.dirname(os.path.abspath(__file__))

    persist_directory = os.path.join(script_dir, 'lance_db')

    if not os.path.exists(persist_directory):
        return None

    try:
        query_text = english_desc
        
        results = query_from_db(
            persist_directory=persist_directory,
            collection_name='product_library',
            query_text=query_text,
            n_results=10,
            strict_match_fields=['制造企业'] if manufacturer else None
        )
        
        if results and english_desc:
            results = reorder_results_with_keyword_match(
                results, english_desc, ['英文描述']
            )

        if results and results.get('ids') and len(results['ids'][0]) > 0:
            best_meta = None
            best_distance = float('inf')

            for i in range(len(results['ids'][0])):
                distance = results['distances'][0][i] if results.get('distances') and len(results['distances'][0]) > i else None
                meta = results['metadatas'][0][i] if results.get('metadatas') and len(results['metadatas'][0]) > i else None

                if distance is not None and distance < 1.2 and distance < best_distance:
                    best_distance = distance
                    best_meta = meta

            if best_meta:
                return best_meta
            else:
                return None
        else:
            return None

    except Exception as e:
        return None


def is_empty_value(val):
    """
    判断值是否为空
    
    Args:
        val: 待判断的值
    
    Returns:
        是否为空
    """
    if val is None or val == '' or val == [] or val == {}:
        return True
    if isinstance(val, (int, float)) and val == 0:
        return True
    return False


def normalize_string(val):
    """
    规范化字符串
    
    Args:
        val: 字符串值
    
    Returns:
        规范化后的字符串
    """
    if isinstance(val, str):
        val = val.strip()
        val = ' '.join(val.split())
    return val


def get_values_from_field(field_data, guid_to_business_type=None, allowed_business_types=None):
    """
    从字段数据中获取值列表
    
    Args:
        field_data: 字段数据
        guid_to_business_type: GUID 到业务类型的映射
        allowed_business_types: 允许的业务类型列表
    
    Returns:
        值列表
    """
    if not isinstance(field_data, list):
        return []
    
    values = []
    
    for item in field_data:
        if not isinstance(item, dict) or 'value' not in item:
            continue
            
        val = item['value']
        
        if is_empty_value(val):
            continue
        
        val = normalize_string(val)
        
        if guid_to_business_type and allowed_business_types:
            guid = item.get('guid')
            if guid:
                business_type = guid_to_business_type.get(guid)
                if business_type not in allowed_business_types:
                    continue
        
        values.append(val)
    
    return values


def get_value_from_field_by_business_type(field_data, guid_to_business_type, business_type):
    """
    根据业务类型从字段数据中获取值
    
    Args:
        field_data: 字段数据
        guid_to_business_type: GUID 到业务类型的映射
        business_type: 业务类型
    
    Returns:
        对应业务类型的值
    """
    if isinstance(field_data, list):
        for item in field_data:
            if isinstance(item, dict) and 'value' in item:
                guid = item.get('guid')
                if guid and guid_to_business_type.get(guid) == business_type:
                    val = item['value']
                    if val is not None and val != '' and not (isinstance(val, (int, float)) and val == 0):
                        if isinstance(val, str):
                            val = val.strip()
                            val = ' '.join(val.split())
                        return val
    return None


def parse_template_placeholder(text):
    """
    解析模板占位符
    
    Args:
        text: 模板文本
    
    Returns:
        (是否可选, 占位符列表)
    """
    if not text:
        return False, []
    
    is_optional = False
    content = text
    
    if text.startswith('[') and text.endswith(']'):
        is_optional = True
        content = text[1:-1]
    
    placeholders = PLACEHOLDER_PATTERN.findall(content)
    
    return is_optional, placeholders


def get_top_left_cell(sheet, row, col):
    """
    获取合并单元格的左上角单元格
    
    Args:
        sheet: 工作表
        row: 行号
        col: 列号
    
    Returns:
        左上角单元格对象
    """
    for merged_range in sheet.merged_cells.ranges:
        if row >= merged_range.min_row and row <= merged_range.max_row and \
           col >= merged_range.min_col and col <= merged_range.max_col:
            return sheet.cell(row=merged_range.min_row, column=merged_range.min_col)
    
    return sheet.cell(row=row, column=col)


def fill_cell_with_values(sheet, row, col, values, original_text=None, placeholder=None, guid_to_business_type=None, helper_sheet=None, helper_row=None):
    """
    用值填充单元格
    
    Args:
        sheet: 工作表
        row: 行号
        col: 列号
        values: 值列表
        original_text: 原始文本
        placeholder: 占位符
        guid_to_business_type: GUID 到业务类型的映射
        helper_sheet: 辅助工作表（用于下拉选项）
        helper_row: 辅助工作表的行号
    
    Returns:
        更新后的 helper_row
    """
    if not OPENPYXL_AVAILABLE:
        return helper_row
    
    cell = get_top_left_cell(sheet, row, col)
    
    seen = set()
    unique_values = []
    
    def make_hashable(v):
        if isinstance(v, list):
            return tuple(v)
        return v
    
    for v in values:
        hashable_v = make_hashable(v)
        if hashable_v not in seen:
            seen.add(hashable_v)
            unique_values.append(v)
    
    font_color = "000000"
    
    source_count = len(values) if values else 0
    
    # 根据来源数量设置颜色：黑色=多个来源一致，红色=多个来源不一致，蓝色=单个来源
    if source_count > 1:
        if len(unique_values) == 1:
            font_color = "000000"
        else:
            font_color = "FF0000"
    elif source_count == 1:
        font_color = "0000FF"
    
    def replace_placeholders_smart(text, placeholder_values):
        if not text:
            return text
        
        is_wrapped_in_brackets = text.startswith('[') and text.endswith(']')
        inner_text = text[1:-1] if is_wrapped_in_brackets else text
        
        all_placeholders = PLACEHOLDER_PATTERN.findall(inner_text)
        
        has_any_value = any(
            ph in placeholder_values and placeholder_values[ph] is not None and placeholder_values[ph] != ""
            for ph in all_placeholders
        )
        
        if is_wrapped_in_brackets and not has_any_value:
            return ""
        
        result = inner_text
        
        for ph, val in placeholder_values.items():
            if val is not None and val != "":
                result = result.replace(f"{{{ph}}}", str(val))
        
        result = PLACEHOLDER_PATTERN.sub('', result)
        
        result = MULTI_PIPE_PATTERN.sub('|', result)
        result = PIPE_START_PATTERN.sub('', result)
        result = PIPE_END_PATTERN.sub('', result)
        
        result = MULTI_SLASH_PATTERN.sub('/', result)
        
        if is_wrapped_in_brackets:
            result = result.strip()
            return result if result else ""
        
        return result
    
    original_font = cell.font
    new_font = None
    if original_font:
        new_font = Font(
            name=original_font.name,
            size=original_font.size,
            bold=original_font.bold,
            italic=original_font.italic,
            vertAlign=original_font.vertAlign,
            underline=original_font.underline,
            strike=original_font.strike,
            color=font_color
        )
    else:
        new_font = Font(color=font_color)
    
    if len(unique_values) == 1:
        if original_text and placeholder:
            placeholder_values = {placeholder: unique_values[0]}
            cell.value = replace_placeholders_smart(original_text, placeholder_values)
        else:
            cell.value = unique_values[0]
        
        cell.font = new_font
    else:
        if original_text and placeholder:
            placeholder_values = {placeholder: unique_values[0] if unique_values else ""}
            cell.value = replace_placeholders_smart(original_text, placeholder_values)
        else:
            cell.value = unique_values[0] if unique_values else ""
        
        cell.font = new_font
        
        # 如果有多个唯一值，创建下拉选择框
        try:
            if helper_sheet is not None and helper_row is not None:
                start_col = 1
                for i, v in enumerate(unique_values):
                    helper_sheet.cell(row=helper_row, column=start_col + i).value = str(v)
                
                end_col = start_col + len(unique_values) - 1
                range_str = f"'{helper_sheet.title}'!${openpyxl.utils.get_column_letter(start_col)}${helper_row}:${openpyxl.utils.get_column_letter(end_col)}${helper_row}"
                
                dv = DataValidation(type="list", formula1=range_str, allow_blank=True)
                sheet.add_data_validation(dv)
                dv.add(cell)
                
                return helper_row + 1
        except Exception as e:
            pass
    
    return helper_row


def clear_cell_placeholder(sheet, row, col, placeholder):
    """
    清除单元格中的占位符
    
    Args:
        sheet: 工作表
        row: 行号
        col: 列号
        placeholder: 占位符文本
    """
    cell = get_top_left_cell(sheet, row, col)
    if cell.value and str(cell.value) == placeholder:
        cell.value = ""


def fill_field_to_cell(sheet, row, col, field_data, guid_to_business_type, helper_sheet, helper_row, use_sum=False):
    """
    将字段数据填充到单元格
    
    Args:
        sheet: 工作表
        row: 行号
        col: 列号
        field_data: 字段数据
        guid_to_business_type: GUID 到业务类型的映射
        helper_sheet: 辅助工作表
        helper_row: 辅助工作表行号
        use_sum: 是否求和
    
    Returns:
        更新后的 helper_row
    """
    if use_sum and isinstance(field_data, list):
        values = get_qty_sum_by_guid(field_data, guid_to_business_type)
    else:
        values = get_values_from_field(field_data, guid_to_business_type) if isinstance(field_data, list) else []
    
    if values:
        return fill_cell_with_values(sheet, row, col, values, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
    return None


def get_qty_sum_by_guid(qty_values, guid_to_business_type):
    """
    按 GUID 分组计算数量总和
    
    Args:
        qty_values: 数量值列表
        guid_to_business_type: GUID 到业务类型的映射
    
    Returns:
        按业务类型分组的总和列表
    """
    business_type_sums = defaultdict(float)
    
    for item in qty_values:
        guid = item.get('guid')
        value = item.get('value')
        
        if guid and guid in guid_to_business_type:
            business_type = guid_to_business_type[guid]
            if value is not None:
                try:
                    business_type_sums[business_type] += float(value)
                except (ValueError, TypeError):
                    pass
    
    sums = list(business_type_sums.values())
    
    if not sums:
        return [item.get('value') for item in qty_values]
    
    return sums


def fill_item_to_rows(sheet, item, start_row, guid_to_business_type, manufacturer=None, helper_sheet=None, helper_row=None):
    """
    将商品项填充到多行（一个商品占用3行）
    
    Args:
        sheet: 工作表
        item: 商品项数据
        start_row: 起始行号
        guid_to_business_type: GUID 到业务类型的映射
        manufacturer: 制造商
        helper_sheet: 辅助工作表
        helper_row: 辅助工作表行号
    
    Returns:
        更新后的 helper_row
    """
    # 查询产品信息
    english_desc = get_value_from_field_by_business_type(item.get('Description', []), guid_to_business_type, '发票')
    
    vector_match = None
    if manufacturer and english_desc:
        vector_match = query_product_info(manufacturer, english_desc)
    
    # 第一行
    row1 = start_row
    
    # GNo - 商品序号
    gno_value = get_value_from_field_by_business_type(item.get('GNo', []), guid_to_business_type, '发票')
    if gno_value:
        new_helper_row = fill_cell_with_values(sheet, row1, 2, [gno_value], guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        clear_cell_placeholder(sheet, row1, 2, '{GNo}')
    
    # 税号 - 从向量数据库获取
    cell = get_top_left_cell(sheet, row1, 3)
    if cell.value and str(cell.value) == '<税号>':
        cell.value = vector_match.get('税号', '') if vector_match else ''
    
    # 品名 - 从向量数据库获取
    cell = get_top_left_cell(sheet, row1, 5)
    if cell.value and str(cell.value) == '<品名>':
        cell.value = vector_match.get('品名', '') if vector_match else ''
    
    # 商品描述
    new_helper_row = fill_field_to_cell(sheet, row1, 11, item.get('Description', []), guid_to_business_type, helper_sheet, helper_row)
    if new_helper_row is not None:
        helper_row = new_helper_row
    else:
        clear_cell_placeholder(sheet, row1, 11, '{Description}')
    
    # 数量1
    new_helper_row = fill_field_to_cell(sheet, row1, 16, item.get('Qty1', []), guid_to_business_type, helper_sheet, helper_row, use_sum=True)
    if new_helper_row is not None:
        helper_row = new_helper_row
    else:
        clear_cell_placeholder(sheet, row1, 16, '{Qty1}')
    
    # 单位1
    new_helper_row = fill_field_to_cell(sheet, row1, 19, item.get('Unit1', []), guid_to_business_type, helper_sheet, helper_row)
    if new_helper_row is not None:
        helper_row = new_helper_row
    else:
        clear_cell_placeholder(sheet, row1, 19, '{Unit1}')
    
    # 申报总价
    new_helper_row = fill_field_to_cell(sheet, row1, 20, item.get('DeclTotal', []), guid_to_business_type, helper_sheet, helper_row)
    if new_helper_row is not None:
        helper_row = new_helper_row
    else:
        clear_cell_placeholder(sheet, row1, 20, '{DeclTotal}')
    
    # 最终目的国
    dest_country_values = get_values_from_field(item.get('DestinationCountry', []))
    if dest_country_values:
        new_helper_row = fill_cell_with_values(sheet, row1, 25, dest_country_values, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        cell = get_top_left_cell(sheet, row1, 25)
        if cell.value and str(cell.value) == '{DestinationCountry}':
            cell.value = ''
    
    # 原产国
    origin_country_values = get_values_from_field(item.get('OriginCountry', []))
    if origin_country_values:
        new_helper_row = fill_cell_with_values(sheet, row1, 23, origin_country_values, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        cell = get_top_left_cell(sheet, row1, 23)
        if cell.value and str(cell.value) == '{OriginCountry}':
            cell.value = ''
    
    # 境内目的地
    district_values = get_values_from_field(item.get('District', []))
    if district_values:
        new_helper_row = fill_cell_with_values(sheet, row1, 30, district_values, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        cell = get_top_left_cell(sheet, row1, 30)
        if cell.value and str(cell.value) == '{District}':
            cell.value = ''
    
    # 征减免税方式
    duty_mode_values = get_values_from_field(item.get('DutyMode', []))
    if duty_mode_values:
        new_helper_row = fill_cell_with_values(sheet, row1, 35, duty_mode_values, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        cell = get_top_left_cell(sheet, row1, 35)
        if cell.value and str(cell.value) == '{DutyMode}':
            cell.value = ''
    
    # 第二行
    row2 = start_row + 1
    
    # 备案序号
    contr_item_values = get_values_from_field(item.get('ContrItem', []))
    if contr_item_values:
        new_helper_row = fill_cell_with_values(sheet, row2, 2, [f"({v})" for v in contr_item_values], guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        cell = get_top_left_cell(sheet, row2, 2)
        if cell.value and str(cell.value) == '(ContrItem)':
            cell.value = ''
    
    # 申报要素 - 从向量数据库获取
    cell = get_top_left_cell(sheet, row2, 5)
    if cell.value and str(cell.value) == '<申报要素>':
        if vector_match and vector_match.get('申报要素'):
            cell.value = vector_match['申报要素']
        else:
            cell.value = ''
    
    # 数量2 或 净重
    qty2_values = item.get('Qty2', [])
    if qty2_values:
        sums = get_qty_sum_by_guid(qty2_values, guid_to_business_type)
        new_helper_row = fill_cell_with_values(sheet, row2, 16, sums, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        netwt_values = get_values_from_field(item.get('NetWt', []))
        if netwt_values:
            new_helper_row = fill_cell_with_values(sheet, row2, 16, netwt_values, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
            if new_helper_row is not None:
                helper_row = new_helper_row
        else:
            cell = get_top_left_cell(sheet, row2, 16)
            if cell.value and (str(cell.value) == '{Qty2}' or str(cell.value) == '{NetWt}'):
                cell.value = ''
    
    # 单位2
    unit2_values = get_values_from_field(item.get('Unit2', []))
    if unit2_values:
        new_helper_row = fill_cell_with_values(sheet, row2, 19, unit2_values, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        cell = get_top_left_cell(sheet, row2, 19)
        if cell.value and str(cell.value) == '{Unit2}':
            cell.value = ''
    
    # 成交币制
    trade_curr_values = get_values_from_field(item.get('TradeCurr', []))
    if trade_curr_values:
        new_helper_row = fill_cell_with_values(sheet, row2, 20, trade_curr_values, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        cell = get_top_left_cell(sheet, row2, 20)
        if cell.value and str(cell.value) == '{TradeCurr}':
            cell.value = ''
    
    # 第三行
    row3 = start_row + 2
    
    # 数量3
    qty3_values = item.get('Qty3', [])
    if qty3_values:
        sums = get_qty_sum_by_guid(qty3_values, guid_to_business_type)
        new_helper_row = fill_cell_with_values(sheet, row3, 16, sums, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        cell = get_top_left_cell(sheet, row3, 16)
        if cell.value and str(cell.value) == '{Qty3}':
            cell.value = ''
    
    # 单位3
    unit3_values = get_values_from_field(item.get('Unit3', []))
    if unit3_values:
        new_helper_row = fill_cell_with_values(sheet, row3, 19, unit3_values, guid_to_business_type=guid_to_business_type, helper_sheet=helper_sheet, helper_row=helper_row)
        if new_helper_row is not None:
            helper_row = new_helper_row
    else:
        cell = get_top_left_cell(sheet, row3, 19)
        if cell.value and str(cell.value) == '{Unit3}':
            cell.value = ''
    
    return helper_row


def fill_declaration_excel(template_path, merged_data, output_dir, guid_to_business_type=None):
    """
    填充报关单 Excel 模板
    
    Args:
        template_path: 模板路径
        merged_data: 合并后的数据
        output_dir: 输出目录
        guid_to_business_type: GUID 到业务类型的映射
    
    Returns:
        输出文件路径或 None
    """
    if not OPENPYXL_AVAILABLE:
        print("错误：openpyxl 未安装，无法生成 Excel 文件")
        return None

    if not os.path.exists(template_path):
        print(f"错误：模板文件不存在：{template_path}")
        return None

    print(f"正在创建输出目录：{output_dir}")
    os.makedirs(output_dir, exist_ok=True)

    try:
        print(f"正在加载模板文件：{template_path}")
        wb = openpyxl.load_workbook(template_path)
        
        # 创建辅助工作表用于存储下拉选项
        helper_sheet = wb.create_sheet(title="_dropdown_options")
        helper_sheet.sheet_state = "hidden"
        helper_row = 1
        
        for sheet in wb.worksheets:
            if sheet.title == "_dropdown_options":
                continue
            
            print(f"正在处理工作表：{sheet.title}")
            
            # 获取制造商信息
            manufacturer = None
            trade_name_values = get_value_from_merged_data(merged_data, "DecMessage.DecHead.TradeName")
            if trade_name_values:
                for item in trade_name_values:
                    value = item.get('value')
                    guid = item.get('guid')
                    if guid and guid_to_business_type.get(guid) == '发票' and value is not None and value != '':
                        manufacturer = value
                        break
            
            processed_cells = set()
            
            # 填充表头字段（前22行）
            for row_idx in range(1, 23):
                for col_idx in range(1, sheet.max_column + 1):
                    top_left = get_top_left_cell(sheet, row_idx, col_idx)
                    cell_key = (top_left.row, top_left.column)
                    
                    if cell_key in processed_cells:
                        continue
                    
                    if not top_left.value or not isinstance(top_left.value, str):
                        continue
                    
                    cell_value = str(top_left.value)
                    is_optional, placeholders = parse_template_placeholder(cell_value)
                    
                    if not placeholders:
                        continue
                    
                    processed_cells.add(cell_key)
                    
                    placeholder_values = {}
                    has_any_value = False
                    
                    for placeholder in placeholders:
                        full_field_path = FIELD_PATH_MAP.get(placeholder, placeholder)
                        values = get_value_from_merged_data(merged_data, full_field_path)
                        
                        field_values = get_values_from_field(values)
                        
                        if field_values:
                            placeholder_values[placeholder] = (field_values[0], field_values)
                            has_any_value = True
                        else:
                            placeholder_values[placeholder] = (None, [])
                    
                    placeholder_first_values = {ph: val[0] for ph, val in placeholder_values.items()}
                    
                    is_wrapped_in_brackets = cell_value.startswith('[') and cell_value.endswith(']')
                    
                    def smart_replace():
                        if not cell_value:
                            return ""
                        
                        inner_text = cell_value[1:-1] if is_wrapped_in_brackets else cell_value
                        
                        all_empty = not any(
                            placeholder_first_values.get(ph) is not None and placeholder_first_values.get(ph) != ""
                            for ph in placeholders
                        )
                        
                        if is_wrapped_in_brackets and all_empty:
                            return ""
                        
                        result = inner_text
                        
                        for ph, val in placeholder_first_values.items():
                            if val is not None and val != "":
                                result = result.replace(f"{{{ph}}}", str(val))
                        
                        result = PLACEHOLDER_PATTERN.sub('', result)
                        
                        result = MULTI_PIPE_PATTERN.sub('|', result)
                        result = PIPE_START_PATTERN.sub('', result)
                        result = PIPE_END_PATTERN.sub('', result)
                        
                        result = MULTI_SLASH_PATTERN.sub('/', result)
                        
                        if is_wrapped_in_brackets:
                            result = result.strip()
                            return result if result else ""
                        
                        return result
                    
                    top_left.value = smart_replace()
                    
                    if has_any_value:
                        first_ph_with_value = None
                        first_field_values = []
                        for placeholder in placeholders:
                            if placeholder_values[placeholder][0] is not None:
                                first_ph_with_value = placeholder
                                first_field_values = placeholder_values[placeholder][1]
                                break
                        
                        if first_ph_with_value:
                            original_font = top_left.font
                            
                            source_count = len(first_field_values) if first_field_values else 0
                            unique_values = []
                            seen = set()
                            for v in first_field_values:
                                if isinstance(v, list):
                                    v_tuple = tuple(v)
                                    if v_tuple not in seen:
                                        seen.add(v_tuple)
                                        unique_values.append(v)
                                else:
                                    if v not in seen:
                                        seen.add(v)
                                        unique_values.append(v)
                            
                            font_color = "000000"
                            if source_count > 1:
                                if len(unique_values) == 1:
                                    font_color = "000000"
                                else:
                                    font_color = "FF0000"
                            elif source_count == 1:
                                font_color = "0000FF"
                            
                            new_font = None
                            if original_font:
                                new_font = Font(
                                    name=original_font.name,
                                    size=original_font.size,
                                    bold=original_font.bold,
                                    italic=original_font.italic,
                                    vertAlign=original_font.vertAlign,
                                    underline=original_font.underline,
                                    strike=original_font.strike,
                                    color=font_color
                                )
                            else:
                                new_font = Font(color=font_color)
                            
                            top_left.font = new_font
                            
                            if len(unique_values) > 1 and helper_sheet is not None and helper_row is not None:
                                try:
                                    start_col = 1
                                    for i, v in enumerate(unique_values):
                                        helper_sheet.cell(row=helper_row, column=start_col + i).value = str(v)
                                    
                                    end_col = start_col + len(unique_values) - 1
                                    range_str = f"'{helper_sheet.title}'!${openpyxl.utils.get_column_letter(start_col)}${helper_row}:${openpyxl.utils.get_column_letter(end_col)}${helper_row}"
                                    
                                    dv = DataValidation(type="list", formula1=range_str, allow_blank=True)
                                    sheet.add_data_validation(dv)
                                    dv.add(top_left)
                                    
                                    helper_row += 1
                                except Exception as e:
                                    pass
            
            # 填充商品列表
            dec_list_data = []
            if 'DecMessage' in merged_data and \
               'DecLists' in merged_data['DecMessage'] and \
               'DecList' in merged_data['DecMessage']['DecLists']:
                dec_list_data = merged_data['DecMessage']['DecLists']['DecList']
            
            if dec_list_data:
                print(f"正在填充 {len(dec_list_data)} 个商品项")
                for item_idx, item in enumerate(dec_list_data):
                    base_row = 23 + 3 * item_idx
                    helper_row = fill_item_to_rows(
                        sheet, item, base_row, guid_to_business_type, manufacturer,
                        helper_sheet=helper_sheet, helper_row=helper_row
                    )
            
            # 删除多余的商品行
            special_relation_row = None
            max_col = sheet.max_column
            for row_idx in range(1, sheet.max_row + 1):
                for col_idx in range(1, max_col + 1):
                    cell = sheet.cell(row=row_idx, column=col_idx)
                    if cell.value and "特殊关系确认" in str(cell.value):
                        special_relation_row = row_idx
                        break
                if special_relation_row:
                    break
            
            if dec_list_data and special_relation_row:
                start_delete_row = 23 + 3 * len(dec_list_data)
                
                if start_delete_row < special_relation_row:
                    rows_to_delete = special_relation_row - start_delete_row
                    print(f"删除 {rows_to_delete} 个多余行")
                    
                    merged_ranges_to_remove = []
                    for merged_range in sheet.merged_cells.ranges:
                        if (merged_range.min_row >= start_delete_row and merged_range.max_row < special_relation_row) or \
                           (merged_range.max_row >= start_delete_row and merged_range.min_row < special_relation_row):
                            merged_ranges_to_remove.append(merged_range)
                    
                    for merged_range in merged_ranges_to_remove:
                        sheet.unmerge_cells(str(merged_range))
                    
                    for i in range(rows_to_delete):
                        sheet.delete_rows(special_relation_row - 1 - i)
        
        # 保存文件
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        json_output_path = os.path.join(output_dir, f"declaration_merged_{timestamp}.json")
        print(f"正在保存合并数据：{json_output_path}")
        with open(json_output_path, 'w', encoding='utf-8') as f:
            json.dump(merged_data, f, ensure_ascii=False, indent=2)
        
        excel_output_path = os.path.join(output_dir, f"报关单_{timestamp}.xlsx")
        print(f"正在保存报关单：{excel_output_path}")
        wb.save(excel_output_path)
        
        return excel_output_path
    except Exception as e:
        print(f"生成 Excel 文件时出错：{e}")
        return None


def main():
    """
    主函数：加载分析结果，合并数据，生成报关单
    """
    print("=" * 60)
    print("报关单生成器")
    print("=" * 60)
    
    parser = argparse.ArgumentParser(description="报关单生成器")
    parser.add_argument(
        "--analysis-results",
        help="分析结果 JSON 文件路径",
        required=True
    )
    parser.add_argument(
        "--output-dir",
        default="/mnt/user-data/outputs",
        help="输出目录（默认：/mnt/user-data/outputs）"
    )
    parser.add_argument(
        "--template-path",
        help="报关单模板路径"
    )

    args = parser.parse_args()
    
    # 如果没有指定模板路径，使用默认模板
    if not args.template_path:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        args.template_path = os.path.join(script_dir, 'assets', '报关单.xlsx')
    
    print(f"分析结果文件：{args.analysis_results}")
    print(f"输出目录：{args.output_dir}")
    print(f"模板文件：{args.template_path}")
    print()

    # 加载分析结果
    print("正在加载分析结果...")
    analysis_results = load_analysis_results(args.analysis_results)
    if not analysis_results:
        print("错误：无法加载分析结果文件")
        return 1

    # 验证数据格式
    print("正在验证数据格式...")
    if not isinstance(analysis_results, list):
        print("错误：分析结果必须是数组格式")
        return 1
    
    for i, result in enumerate(analysis_results):
        if not isinstance(result, dict):
            print(f"错误：第 {i} 个元素不是字典")
            return 1
        if 'guid' not in result:
            print(f"错误：第 {i} 个元素缺少 'guid' 字段")
            return 1
        if 'analysis' not in result:
            print(f"错误：第 {i} 个元素缺少 'analysis' 字段")
            return 1
        if not isinstance(result['analysis'], dict):
            print(f"错误：第 {i} 个元素的 'analysis' 字段不是字典")
            return 1
    
    print(f"验证通过，共 {len(analysis_results)} 个单证")
    print()

    # 合并数据
    print("正在合并数据...")
    merged_data, guid_to_business_type = merge_json_objects_with_guid(analysis_results)
    print("数据合并完成")
    print()

    # 查询产品信息
    if 'DecMessage' in merged_data and 'DecLists' in merged_data['DecMessage']:
        print("正在查询产品信息...")
        dec_list = merged_data['DecMessage']['DecLists'].get('DecList', [])
        for item in dec_list:
            manufacturer = None
            english_desc = None

            if 'DecHead' in merged_data.get('DecMessage', {}):
                trade_name = get_value_from_merged_data(merged_data, 'DecMessage.DecHead.TradeName')
                if trade_name and len(trade_name) > 0:
                    for item_guid in trade_name:
                        guid = item_guid.get('guid')
                        if guid and guid_to_business_type.get(guid) == '发票':
                            manufacturer = item_guid.get('value')
                            break

            desc_list = item.get('Description', [])
            if desc_list and len(desc_list) > 0:
                for item_guid in desc_list:
                    guid = item_guid.get('guid')
                    if guid and guid_to_business_type.get(guid) == '发票':
                        english_desc = item_guid.get('value')
                        break

            if english_desc:
                product_info = query_product_info(manufacturer, english_desc)
                if product_info:
                    print(f"  找到产品：{product_info.get('品名', '未知')}")
        print()

    # 生成 Excel 文件
    print("正在生成报关单 Excel...")
    excel_path = fill_declaration_excel(
        args.template_path,
        merged_data,
        args.output_dir,
        guid_to_business_type
    )

    if excel_path:
        print()
        print("=" * 60)
        print(f"报关单生成成功！")
        print(f"文件路径：{excel_path}")
        print("=" * 60)
        return 0
    else:
        print()
        print("=" * 60)
        print("报关单生成失败！")
        print("=" * 60)
        return 1


if __name__ == "__main__":
    sys.exit(main())
